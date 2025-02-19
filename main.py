from fastapi import FastAPI, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
import pandas as pd
import sqlalchemy
import io
import openpyxl
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, PieChart

app = FastAPI()

# Configuración de CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["https://api-familia-tareas-node.onrender.com"],  
    allow_credentials=True,
    allow_methods=["GET"],
    allow_headers=["*"],
)

# Conexión a la base de datos
DATABASE_URL = "postgresql://neondb_owner:npg_W1qfEuBdANj4@ep-square-block-a8x6hhxv-pooler.eastus2.azure.neon.tech/neondb?sslmode=require"
engine = sqlalchemy.create_engine(DATABASE_URL)

@app.get("/export/finances")
def exportar_finanzas(year: int = Query(None), month: int = Query(None)):
    """Genera un archivo Excel con los gastos, ingresos y análisis financiero con gráficos."""
    try:
        # Consultas SQL
        expenses_query = f"""
            SELECT e.fecha, e.valor, e.description, c.name AS categoria, a.name AS cuenta, u.name AS usuario
            FROM "Expenses" e
            LEFT JOIN "Categories" c ON e.categoria_id = c.id
            LEFT JOIN "Accounts" a ON e.cuenta_id = a.id
            LEFT JOIN "Users" u ON e.user_id = u.id
            WHERE EXTRACT(YEAR FROM e.fecha) = {year} AND EXTRACT(MONTH FROM e.fecha) = {month}
        """

        incomes_query = f"""
            SELECT i.fecha, i.valor, i.description, c.name AS categoria, a.name AS cuenta, u.name AS usuario
            FROM "Incomes" i
            LEFT JOIN "Categories" c ON i.categoria_id = c.id
            LEFT JOIN "Accounts" a ON i.cuenta_id = a.id
            LEFT JOIN "Users" u ON i.user_id = u.id
            WHERE EXTRACT(YEAR FROM i.fecha) = {year} AND EXTRACT(MONTH FROM i.fecha) = {month}
        """

        # Cargar datos
        expenses_df = pd.read_sql(expenses_query, con=engine)
        incomes_df = pd.read_sql(incomes_query, con=engine)

        # Cálculos financieros
        total_expenses = expenses_df["valor"].sum() if not expenses_df.empty else 0
        total_incomes = incomes_df["valor"].sum() if not incomes_df.empty else 0
        balance = total_incomes - total_expenses
        porcentaje_gastado = (total_expenses / total_incomes * 100) if total_incomes > 0 else 0

        category_spending = expenses_df.groupby("categoria")["valor"].sum().reset_index()
        category_spending["% del Gasto"] = (category_spending["valor"] / total_expenses * 100).round(2)

        # Crear archivo Excel
        output = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "Análisis Financiero"

        # Escribir datos en Excel
        ws.append(["Métrica", "Valor"])
        ws.append(["Total Ingresos", total_incomes])
        ws.append(["Total Gastos", total_expenses])
        ws.append(["Balance", balance])
        ws.append(["Porcentaje de Ingresos Gastados", f"{porcentaje_gastado:.2f}%"])

        # Agregar gráficos
        chart = BarChart()
        data = Reference(ws, min_col=2, min_row=2, max_row=5)
        categories = Reference(ws, min_col=1, min_row=2, max_row=5)
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(categories)
        chart.title = "Análisis Financiero"
        ws.add_chart(chart, "E5")

        ws2 = wb.create_sheet("Gastos por Categoría")
        ws2.append(["Categoría", "Valor", "% del Gasto"])
        for row in category_spending.itertuples():
            ws2.append([row.categoria, row.valor, f"{row._3}%"])

        pie_chart = PieChart()
        labels = Reference(ws2, min_col=1, min_row=2, max_row=len(category_spending) + 1)
        data = Reference(ws2, min_col=2, min_row=2, max_row=len(category_spending) + 1)
        pie_chart.add_data(data, titles_from_data=False)
        pie_chart.set_categories(labels)
        pie_chart.title = "Distribución de Gastos"
        ws2.add_chart(pie_chart, "E5")

        wb.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=finanzas_familia.xlsx"}
        )
    except Exception as e:
        return {"error": str(e)}
